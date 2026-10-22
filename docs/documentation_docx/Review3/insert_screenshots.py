"""Insert the verified capture inventory into the existing Review 3 DOCX."""
import json
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[3]
folder = Path(__file__).resolve().parent
target = folder / "Review3.docx"
backup = ROOT / ".git/review3-original.docx"
if not backup.exists():
    copy2(target, backup)
doc = Document(backup)
captures = json.loads((folder / "screenshots/captures.json").read_text(encoding="utf-8"))
screen_heading = next(p for p in doc.paragraphs if p.text.strip() == "Screen shots")
start = doc.paragraphs.index(screen_heading) if screen_heading in doc.paragraphs else next(i for i,p in enumerate(doc.paragraphs) if p.text.strip() == "Screen shots")
anchor = next(p for p in doc.paragraphs[start+1:] if p._p.find(".//"+qn("w:sectPr")) is not None)


def add(text="", style=None):
    paragraph = anchor.insert_paragraph_before(text, style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


add("Captured on 22 October 2026 from the refreshed source candidate based on commit 14f6a414. "
    "The web and Flutter web clients use the real Gateway, six packaged business applications, "
    "PostgreSQL and RabbitMQ in an isolated demonstration environment. Patient and staff details "
    "are synthetic. Screenshots show the observed runtime state.")
add("Mobile GUI evidence uses Flutter web in a phone viewport. It does not establish native "
    "Android or iOS execution. Unavailable workspaces and report coverage restrictions are "
    "identified explicitly; inpatient care and surgery remain outside this review scope.")
add("Evidence coverage")
for text in [
    "1. Shared web pages include landing, login and access-denied states.",
    "2. Web roles include Administrator, Manager, Doctor, Nurse, Laboratory technician, Pharmacist, Cashier and Patient. Shared workflows are illustrated under the role that performs them.",
    "3. Mobile evidence includes login, all eight role workspaces, organization, patients, appointment filtering, laboratory filtering, Billing payment and notification details.",
    "4. Recorded limitations include Lab start rejection, missing outpatient report reconciliation, the legacy invoice view, and unfinished mobile workspaces. These are observed limitations, not completed functionality."]:
    add(text)
order = ["Shared web screens", "Administrator — web", "Manager — web", "Doctor — web", "Nurse — web", "Laboratory technician — web", "Pharmacist — web", "Cashier — web", "Patient — web", "Mobile GUI — shared", "Mobile GUI — administrator", "Mobile GUI — manager", "Mobile GUI — doctor", "Mobile GUI — nurse", "Mobile GUI — laboratory technician", "Mobile GUI — pharmacist", "Mobile GUI — cashier", "Mobile GUI — patient"]
captures.sort(key=lambda c: (order.index(c["group"]), c["file"]))
(folder / "screenshots/captures.json").write_text(json.dumps(captures, ensure_ascii=False, indent=2), encoding="utf-8")
previous_group = None
group_number = 0
item_number = 0
for capture in captures:
    new_group = capture["group"] != previous_group
    if new_group:
        group_number += 1
        item_number = 0
        previous_group = capture["group"]
        paragraph = add(f"{group_number}. {previous_group.replace(" — ", " ")}", "Heading 2")
        paragraph.paragraph_format.page_break_before = True
        for run in paragraph.runs: run.font.color.rgb = RGBColor(0, 0, 0)
        paragraph.paragraph_format.keep_with_next = True
    item_number += 1
    paragraph = add(f"{group_number}.{item_number} {capture['caption']}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.runs[0].bold = True
    if not new_group: paragraph.paragraph_format.page_break_before = True
    paragraph = add()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_path = folder / capture["file"]
    width_px, height_px = Image.open(image_path).size
    width = min(2.6 if capture["platform"] == "mobile" else 6.5, 7.8 * width_px / height_px)
    picture = paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    paragraph.paragraph_format.keep_with_next = bool(capture.get("note"))
    picture._inline.docPr.set("descr", capture["caption"])
    if capture.get("note"):
        paragraph = add(capture["note"])
        for run in paragraph.runs:
            run.font.size = Pt(10)
doc.save(target)
print(f"Inserted {len(captures)} screenshots in {group_number} indexed groups; "
      f"preserved {len(doc.tables)} existing tables without adding screenshot tables.")
